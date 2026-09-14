"""Shared python-docx helpers for every CycloTool Web report.

Every report (Busbar, and each future one -- Water Cooling, Harmonics,
Design Data, Engineering Data) follows the same visual pattern: a title,
a project metadata block, remark/assumption bullets, one or more data
tables with pass/fail shading, and a footer. Centralising that here means
a new report is "call these helpers with this data", not "re-solve docx
formatting from scratch".
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Colors matching the web UI's pass/fail conventions.
COLOR_OK_BG = "E9F7EC"
COLOR_FAIL_BG = "FBEAEA"
COLOR_HEADER_BG = "EEF0F3"
COLOR_MUTED = RGBColor(0x5B, 0x67, 0x73)
COLOR_FAIL_TEXT = RGBColor(0xB3, 0x26, 0x1E)
COLOR_OK_TEXT = RGBColor(0x1C, 0x7A, 0x3E)

FONT_NAME = "Calibri"


@dataclass
class ProjectMeta:
    name: str = "Untitled Project"
    equipment: str = ""
    client: str = ""
    contractor: str = ""
    project_no: str = ""
    customer_po: str = ""
    prepared_by: str = ""
    date: str = ""
    extra: dict[str, str] = field(default_factory=dict)


def new_document(*, landscape: bool = False) -> Document:
    doc = Document()
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(9.5)
    # Ensure east-asian font fallback doesn't override the Latin font name
    # (a common python-docx gotcha where only rFonts/w:ascii gets set).
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)

    return doc


def add_title(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)


def add_meta_table(doc: Document, project: ProjectMeta, extra_rows: list[tuple[str, str]]) -> None:
    rows = [
        ("Equipment", project.equipment),
        ("Project", project.name),
        ("Client", project.client),
        ("Contractor", project.contractor),
        ("ABB-equivalent Project No.", project.project_no),
        ("Customer P.O. No.", project.customer_po),
        ("Prepared by", f"{project.prepared_by} – {project.date}"),
        *extra_rows,
        *project.extra.items(),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.autofit = True
    for label, value in rows:
        if not label and not value:
            continue
        row = table.add_row()
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[0].paragraphs[0].runs[0].font.color.rgb = COLOR_MUTED
        row.cells[1].text = str(value)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = COLOR_MUTED


def add_section_heading(doc: Document, text: str) -> None:
    doc.add_heading(text, level=1)


def _shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color})
    tc_pr.append(shd)


def add_data_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    row_status: list[bool | None] | None = None,
    first_col_left_align: bool = True,
) -> None:
    """A header-shaded data table with optional per-row pass/fail shading.

    row_status[i] = True -> shade row i green, False -> red, None -> no shading.
    """
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        _shade_cell(hdr_cells[i], COLOR_HEADER_BG)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(8.5)

    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        status = row_status[r_idx] if row_status else None
        fill = COLOR_OK_BG if status is True else COLOR_FAIL_BG if status is False else None
        for c_idx, value in enumerate(row):
            cells[c_idx].text = str(value)
            if fill:
                _shade_cell(cells[c_idx], fill)
            align_left = first_col_left_align and c_idx == 0
            for p in cells[c_idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if align_left else WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs:
                    run.font.size = Pt(8.5)

    doc.add_paragraph()


def add_note(doc: Document, text: str, *, color: RGBColor | None = None, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = color or COLOR_MUTED
    run.font.bold = bold


def add_footer_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    run.font.italic = True


def render_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
